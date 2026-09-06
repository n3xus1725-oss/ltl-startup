"""Spreadsheet (CSV/XLSX) table and row extractor."""

import csv
import io
from typing import Any, Dict, List, Tuple

from packages.domain.logging import logger


class SpreadsheetExtractor:
    """Extracts rows and tables from CSV and Excel files."""

    @staticmethod
    def extract_csv(content: bytes) -> Tuple[str, List[List[str]], Dict[str, Any]]:
        text_stream = io.StringIO(content.decode("utf-8", errors="replace"))
        reader = csv.reader(text_stream)
        rows: List[List[str]] = []
        text_lines: List[str] = []

        for r in reader:
            if r:
                rows.append([c.strip() for c in r])
                text_lines.append(" | ".join([c.strip() for c in r if c.strip()]))

        return "\n".join(text_lines), rows, {"row_count": len(rows), "type": "csv"}

    @staticmethod
    def extract_xlsx(content: bytes) -> Tuple[str, List[List[str]], Dict[str, Any]]:
        rows: List[List[str]] = []
        text_lines: List[str] = []
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                for r in sheet.iter_rows(values_only=True):
                    row_vals = [str(cell).strip() for cell in r if cell is not None]
                    if row_vals:
                        rows.append(row_vals)
                        text_lines.append(" | ".join(row_vals))
            wb.close()
        except Exception as e:
            logger.warning(f"Failed to extract Excel workbook: {e}")
            # Fallback to CSV parsing if simple text
            return SpreadsheetExtractor.extract_csv(content)

        return "\n".join(text_lines), rows, {"row_count": len(rows), "type": "xlsx"}
